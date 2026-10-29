"""A5-09..14 — export writers (html/csv/json/txt/xlsx) and the runner.

Two layers:
  * writer layer: stubbed Messages (text/image/file/link/system) -> temp dir,
    then assert on the produced bytes.
  * runner layer: the real decrypted sample drives ``export.runner.run``.

All writes go to a per-test temp dir under ``.tmp`` / ``%TEMP%``.
"""
from __future__ import annotations

import csv
import io
import json
import sys
import unittest
from pathlib import Path
from unittest import mock

from _support import TempDirMixin, make_messages

from wexport.export import csv_export, html_export, json_export, runner, txt_export, xlsx_export

DISPLAY = "张三（备注）"
CHAT = "wxid_friend_x"
CONTACTS = {
    CHAT: {"remark": "张三（备注）", "nick": "张三", "alias": "zs", "type": 1},
    "wxid_self_001": {"remark": "", "nick": "我自己", "alias": "", "type": 0},
}


class WriterBase(TempDirMixin):
    prefix = "export-"

    @classmethod
    def setUpClass(cls):
        cls.msgs = make_messages()

    def assert_exists_nonempty(self, path: Path) -> Path:
        self.assertIsNotNone(path)
        self.assertTrue(path.is_file(), f"{path} 未生成")
        self.assertGreater(path.stat().st_size, 0, f"{path} 是空文件")
        return path


# --------------------------------------------------------------------------- #
# HTML
# --------------------------------------------------------------------------- #
class TestHtmlWriter(WriterBase):
    def test_writes_utf8_html_with_meta_and_text(self):
        out = self.sub("html")
        path = self.assert_exists_nonempty(
            html_export.write_chat(out, CHAT, DISPLAY, self.msgs, CONTACTS))
        self.assertEqual(path.name, f"{DISPLAY}.html")
        raw = path.read_bytes()
        self.assertTrue(raw.startswith(b"<!DOCTYPE html>"))
        self.assertIn(b'<meta charset="utf-8">', raw,
                      'HTML 缺少 UTF-8 meta，中文会乱码')
        text = raw.decode("utf-8")
        self.assertIn("你好，这是文本消息", text)
        self.assertIn("项目需求说明书", text)
        self.assertIn("微信 4.x 数据目录结构解析", text)
        self.assertIn("可以开始聊天了", text)
        # one block per message
        self.assertEqual(text.count('<div class="msg '), len(self.msgs))
        self.assertIn(f"{len(self.msgs)} 条消息", text)

    def test_html_escapes_markup(self):
        out = self.sub("html-esc")
        path = html_export.write_chat(out, CHAT, DISPLAY, self.msgs, CONTACTS)
        text = path.read_text(encoding="utf-8")
        self.assertIn("&lt;world&gt;", text, "尖括号未转义，页面结构会被消息内容破坏")
        self.assertIn("&amp;", text)
        self.assertNotIn("<world>", text)

    def test_image_and_link_and_system_rendering(self):
        out = self.sub("html-kinds")
        text = html_export.write_chat(out, CHAT, DISPLAY, self.msgs, CONTACTS).read_text("utf-8")
        self.assertIn("media-missing", text)          # image without local file
        self.assertIn('class="card"', text)           # link card
        self.assertIn("https://example.com/a?x=1&amp;y=2", text)
        self.assertIn('class="sys"', text)            # system message style

    def test_media_dir_relative_links(self):
        out = self.sub("html-media")
        media = self.sub("html-media/user.png")
        media.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)
        md5 = self.msgs[1].extras["md5"]
        path = html_export.write_chat(out, CHAT, DISPLAY, self.msgs, CONTACTS,
                                      {md5: [media]}, "media/张三")
        text = path.read_text("utf-8")
        self.assertIn('src="media/张三/user.png"', text)

    def test_safe_name_blocks_path_traversal(self):
        out = self.sub("html-safe")
        path = html_export.write_chat(out, CHAT, "../../evil", self.msgs, CONTACTS)
        self.assertEqual(path.parent.resolve(), out.resolve(), "会话名逃出了导出目录")
        self.assertEqual(path.name, "....evil.html")
        for bad in ('a/b', 'a\\b', 'a:b', 'a*b', 'a?b', 'a"b', 'a<b', 'a>b', 'a|b'):
            cleaned = html_export.safe_name(bad)
            self.assertNotIn("/", cleaned)
            self.assertNotIn("\\", cleaned)
            self.assertNotIn(":", cleaned)
        self.assertEqual(html_export.safe_name("   "), "chat")
        self.assertEqual(html_export.safe_name(""), "chat")

    def test_index_page(self):
        out = self.sub("html-index")
        stats = [{"chat": CHAT, "display": DISPLAY, "count": len(self.msgs)},
                 {"chat": "wxid_other", "display": "李四", "count": 3}]
        path = self.assert_exists_nonempty(html_export.write_index(out, stats))
        self.assertEqual(path.name, "index.html")
        text = path.read_text("utf-8")
        self.assertIn("<meta charset=\"utf-8\">", text)
        self.assertIn(f'href="{DISPLAY}.html"', text)
        self.assertIn("2 个会话", text)
        self.assertIn(str(len(self.msgs) + 3), text)


# --------------------------------------------------------------------------- #
# CSV
# --------------------------------------------------------------------------- #
class TestCsvWriter(WriterBase):
    def test_bom_and_header(self):
        out = self.sub("csv")
        path = self.assert_exists_nonempty(csv_export.write_all(out, DISPLAY, self.msgs))
        raw = path.read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"), "CSV 缺少 UTF-8 BOM，Excel 会乱码")
        rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
        self.assertEqual(rows[0], csv_export.HEADERS)
        self.assertEqual(rows[0], ["时间", "发送者", "类型", "内容", "链接/文件", "local_id", "local_type"])
        self.assertEqual(len(rows), len(self.msgs) + 1)
        for row in rows[1:]:
            self.assertEqual(len(row), len(csv_export.HEADERS))

    def test_rows_carry_message_data(self):
        out = self.sub("csv-rows")
        path = csv_export.write_all(out, DISPLAY, self.msgs)
        rows = list(csv.reader(io.StringIO(path.read_text("utf-8-sig"))))[1:]
        kinds = [r[2] for r in rows]
        self.assertEqual(kinds, ["text", "image", "file", "link", "system"])
        self.assertEqual(rows[0][1], "wxid_friend_x")
        self.assertIn("你好，这是文本消息", rows[0][3])
        self.assertEqual(rows[2][4], self.msgs[2].extras["file_name"])
        self.assertEqual(rows[3][4], self.msgs[3].url)
        self.assertEqual(rows[4][6], "0x2710")

    def test_roundtrip_utf8(self):
        out = self.sub("csv-utf8")
        path = csv_export.write_all(out, DISPLAY, self.msgs)
        text = path.read_text(encoding="utf-8-sig")
        self.assertIn("项目需求说明书", text)


# --------------------------------------------------------------------------- #
# JSON
# --------------------------------------------------------------------------- #
class TestJsonWriter(WriterBase):
    def test_deserialisable_and_complete(self):
        out = self.sub("json")
        path = self.assert_exists_nonempty(json_export.write_chat(out, DISPLAY, self.msgs))
        data = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(len(data), len(self.msgs))
        keys = {"time", "time_str", "sender", "kind", "text", "title", "url",
                "local_id", "server_id", "local_type", "extras"}
        for rec in data:
            self.assertEqual(set(rec), keys)
        self.assertEqual([r["kind"] for r in data], ["text", "image", "file", "link", "system"])
        self.assertEqual(data[0]["text"], self.msgs[0].content)

    def test_sensitive_extras_are_stripped(self):
        out = self.sub("json-sec")
        path = json_export.write_chat(out, DISPLAY, self.msgs)
        data = json.loads(path.read_text(encoding="utf-8"))
        for rec in data:
            self.assertNotIn("aeskey", rec["extras"], "导出文件泄露了媒体 aeskey")
            self.assertNotIn("attachid", rec["extras"])

    def test_ensure_ascii_false(self):
        out = self.sub("json-cn")
        path = json_export.write_chat(out, DISPLAY, self.msgs)
        self.assertIn("项目需求说明书", path.read_text("utf-8"))

    def test_empty_messages(self):
        out = self.sub("json-empty")
        path = json_export.write_chat(out, DISPLAY, [])
        self.assertEqual(json.loads(path.read_text("utf-8")), [])


# --------------------------------------------------------------------------- #
# TXT
# --------------------------------------------------------------------------- #
class TestTxtWriter(WriterBase):
    def test_header_and_lines(self):
        out = self.sub("txt")
        path = self.assert_exists_nonempty(txt_export.write_chat(out, DISPLAY, self.msgs))
        text = path.read_text(encoding="utf-8")
        self.assertTrue(text.startswith(f"# {DISPLAY}"))
        lines = [ln for ln in text.splitlines() if ln.startswith("[")]
        self.assertEqual(len(lines), len(self.msgs))
        for msg in self.msgs:
            self.assertIn(msg.display_text, text)
        self.assertIn(f"wxid_friend_x: {self.msgs[0].display_text}", text)

    def test_utf8_no_bom_needed(self):
        out = self.sub("txt-enc")
        path = txt_export.write_chat(out, DISPLAY, self.msgs)
        self.assertIn("你好，这是文本消息", path.read_text("utf-8"))


# --------------------------------------------------------------------------- #
# XLSX
# --------------------------------------------------------------------------- #
class TestXlsxWriter(WriterBase):
    def test_openable_by_openpyxl(self):
        out = self.sub("xlsx")
        path = self.assert_exists_nonempty(xlsx_export.write_chat(out, DISPLAY, self.msgs))
        from openpyxl import load_workbook

        wb = load_workbook(path)
        try:
            ws = wb.active
            self.assertEqual(ws.title, "聊天记录")
            self.assertEqual(ws.max_row, len(self.msgs) + 1)
            self.assertEqual(ws.max_column, 5)
            header = [c.value for c in ws[1]]
            self.assertEqual(header, ["时间", "发送者", "类型", "内容", "链接/文件"])
            self.assertTrue(ws["A1"].font.bold, "表头未加粗")
            self.assertEqual(ws.freeze_panes, "A2")
            self.assertEqual(ws.cell(row=2, column=4).value, self.msgs[0].display_text)
            self.assertEqual(ws.cell(row=4, column=5).value, self.msgs[2].extras["file_name"])
            self.assertEqual(ws.cell(row=5, column=5).value, self.msgs[3].url)
        finally:
            wb.close()

    def test_returns_none_without_openpyxl(self):
        out = self.sub("xlsx-noopenpyxl")
        with mock.patch.dict(sys.modules, {"openpyxl": None}):
            self.assertIsNone(xlsx_export.write_chat(out, DISPLAY, self.msgs))


# --------------------------------------------------------------------------- #
# Runner
# --------------------------------------------------------------------------- #
class TestRunnerRealSample(TempDirMixin, unittest.TestCase):
    """TC-EXP-15: end-to-end export over the real decrypted sample."""

    prefix = "runner-"

    @classmethod
    def setUpClass(cls):
        from _support import DECRYPTED, require_path

        require_path(DECRYPTED / "session" / "session.db", "明文样本 probe/decrypted")
        cls.decrypted = DECRYPTED

    def _run(self, **kw):
        req = runner.ExportRequest(
            decrypted_dir=self.decrypted,
            account_dir=self.decrypted,
            wxid="wxid_test",
            out_dir=self.tmp / "out",
            export_media=False,
            **kw,
        )
        return runner.run(req)

    def test_all_formats_written_for_every_chat(self):
        out = self.tmp / "out"
        res = self._run(formats=["html", "csv", "json", "txt", "xlsx"])
        self.assertIsInstance(res, dict)
        self.assertEqual(Path(res["out_dir"]), out)
        self.assertGreaterEqual(len(res["chats"]), 2, "真实样本应导出至少 2 个会话")
        self.assertGreater(sum(s["count"] for s in res["chats"]), 0)
        self.assertTrue((out / "index.html").is_file())
        self.assertTrue((out / "style.css").is_file(), "HTML 缺少 style.css，样式会丢失")
        for stat in res["chats"]:
            safe = html_export.safe_name(stat["display"])
            for ext in ("html", "csv", "json", "txt", "xlsx"):
                p = out / f"{safe}.{ext}"
                self.assertTrue(p.is_file(), f"{stat['chat']} 缺少 {ext}")
                self.assertGreater(p.stat().st_size, 0)
            data = json.loads((out / f"{safe}.json").read_text("utf-8"))
            self.assertEqual(len(data), stat["count"])

    def test_html_of_real_chat_contains_utf8_meta(self):
        out = self.tmp / "out"
        res = self._run(formats=["html"])
        stat = max(res["chats"], key=lambda s: s["count"])
        safe = html_export.safe_name(stat["display"])
        raw = (out / f"{safe}.html").read_bytes()
        self.assertIn(b'<meta charset="utf-8">', raw)
        self.assertIn(f"{stat['count']} 条消息".encode(), raw)

    def test_chat_subset(self):
        out = self.tmp / "out"
        res = self._run(chats=["filehelper"], formats=["txt"])
        self.assertEqual([s["chat"] for s in res["chats"]], ["filehelper"])
        self.assertFalse((out / "index.html").exists(), "只导出 txt 时不应生成 index.html")
        self.assertTrue(any(out.glob("*.txt")))

    def test_media_disabled_creates_no_media_dir(self):
        out = self.tmp / "out"
        self._run(formats=["html"])
        self.assertFalse((out / "media").exists(), "export_media=False 仍然创建了 media 目录")

    def test_default_request_is_html_only(self):
        req = runner.ExportRequest(decrypted_dir=self.decrypted, account_dir=self.decrypted,
                                   wxid="w", out_dir=self.tmp / "o")
        self.assertEqual(req.formats, ["html"])
        self.assertTrue(req.export_media)
        self.assertTrue(req.include_system)
        self.assertEqual((req.since, req.until), (0, 0))

    def test_unknown_chat_is_skipped(self):
        res = self._run(chats=["no_such_chat_xyz"], formats=["html"])
        self.assertEqual(res["chats"], [])
        self.assertTrue((self.tmp / "out" / "index.html").is_file())


if __name__ == "__main__":
    sys.exit(unittest.main())
